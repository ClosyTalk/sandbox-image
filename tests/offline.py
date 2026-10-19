"""Run inside the image with Docker networking disabled; only loopback is used."""
import http.server
import importlib.util
import json
import os
from pathlib import Path
import platform
import pwd
import shutil
import socket
import subprocess
import tempfile
import threading
import unittest


def run(*args, timeout=45):
    return subprocess.run(args, check=True, capture_output=True, text=True, timeout=timeout).stdout


class OfflineRuntime(unittest.TestCase):
    def test_no_build_tools_in_runtime(self):
        for module in ("pip", "setuptools", "wheel"):
            self.assertIsNone(importlib.util.find_spec(module), f"Do not ship {module}")
        for binary in ("pip", "pip3", "ffprobe", "gcc", "make"):
            self.assertIsNone(shutil.which(binary), f"Do not ship {binary}")
        packages = run("dpkg-query", "-W", "-f=${Package}\n").splitlines()
        for package in ("ffmpeg", "libvpx9", "python3-pip", "python3-venv", "python3-pip-whl", "python3-setuptools-whl"):
            self.assertNotIn(package, packages)

    def test_ffmpeg_as_sandbox_user(self):
        self.assertEqual(os.getuid(), 10001)
        self.assertEqual(pwd.getpwuid(os.getuid()).pw_name, "sandbox")
        self.assertIsNotNone(shutil.which("ffmpeg"), "Recording requires ffmpeg on PATH")
        version = run("ffmpeg", "-version").splitlines()[0]
        self.assertIn("ffmpeg version", version)
        self.assertIn("GNU Lesser General Public", run("ffmpeg", "-L"))
        print(f"\n{version} (user=sandbox)", flush=True)

    def test_browser_recording_offline(self):
        self.check_browser_recording(cursor=False)

    def test_browser_recording_with_cursor_offline(self):
        self.check_browser_recording(cursor=True)

    def check_browser_recording(self, cursor):
        self.assertEqual(os.getuid(), 10001)

        def browser(*args):
            result = subprocess.run(
                ("agent-browser", "--session", f"offline-recording-{cursor}", "--json", *args),
                capture_output=True, text=True, timeout=45,
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            response = json.loads(result.stdout)
            self.assertTrue(response["success"], response)
            return response["data"]

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            page = root / "recording.html"
            page.write_text('''<!doctype html><html><title>Offline recording</title>
            <body><label>Name <input id="name"></label>
            <button onclick="document.querySelector('#result').textContent='Saved '+document.querySelector('#name').value">Save</button>
            <p id="result">Waiting</p></body></html>''')
            video = root / "recording.webm"
            try:
                browser("open", page.as_uri())
                browser("set", "viewport", "801", "601")
                browser("record", "start", str(video), *(("--cursor",) if cursor else ()))
                browser("fill", "#name", "ClosyTalk")
                browser("click", "button")
                self.assertIn("Saved ClosyTalk", json.dumps(browser("get", "text", "#result")))
                browser("record", "stop")
                self.assertTrue(video.is_file(), "Recording must produce a WebM file")
                self.assertGreater(video.stat().st_size, 0)
                # EBML header and DocType, then decode an actual video frame with
                # the bundled Chromium. No ffprobe or extra FFmpeg decoder needed.
                header = video.read_bytes()[:64]
                self.assertTrue(header.startswith(b"\x1a\x45\xdf\xa3"), header)
                self.assertIn(b"\x42\x82\x84webm", header)
                player = root / "player.html"
                player.write_text('''<!doctype html><video muted autoplay src="recording.webm"></video>
                <script>
                const video = document.querySelector('video');
                video.requestVideoFrameCallback((now, frame) => {
                    const result = document.createElement('p');
                    result.id = 'decoded';
                    result.textContent = JSON.stringify({width: video.videoWidth,
                        height: video.videoHeight, frames: frame.presentedFrames});
                    document.body.append(result);
                });
                </script>''')
                browser("open", player.as_uri())
                browser("wait", "#decoded")
                probe = json.loads(browser("get", "text", "#decoded")["text"])
                self.assertEqual(probe["width"], 802, "Pad odd viewport dimensions for VP8")
                self.assertEqual(probe["height"], 602, "Pad odd viewport dimensions for VP8")
                self.assertGreater(probe["frames"], 0)
                print(f"\nOffline recording: {video.stat().st_size} bytes; {json.dumps(probe)}", flush=True)
            finally:
                browser("close")

    def test_runtime_and_no_external_network(self):
        self.assertEqual(os.getuid(), 10001, "Agent code must run as the sandbox user")
        self.assertEqual(platform.python_version_tuple()[:2], ("3", "12"))
        self.assertEqual(run("node", "--version").strip(), "v22.23.3")
        self.assertIn("0.39.0", run("agent-browser", "--version"))
        self.assertIn("156.0.8078.4", run("chromium", "--version"))
        self.assertIsNone(shutil.which("npm"), "Do not ship the build-only npm dependency tree in the runtime")
        with self.assertRaises(OSError):
            socket.create_connection(("1.1.1.1", 443), timeout=1)
        for directory in ("/opt/skill", "/opt/skills", "/tmp/downloads"):
            with tempfile.TemporaryDirectory(dir=directory) as mounted:
                target = Path(mounted) / "resource.bin"
                target.write_bytes(b"\x00\xffsandbox")
                self.assertEqual(target.read_bytes(), b"\x00\xffsandbox")

    def test_documents_round_trip(self):
        from docx import Document
        from docxtpl import DocxTemplate
        from openpyxl import load_workbook
        import pandas as pd
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches
        from reportlab.pdfgen.canvas import Canvas
        from PIL import Image
        import matplotlib.pyplot as plt

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            doc = Document()
            doc.add_paragraph("Customer {{ name }}")
            doc.save(root / "template.docx")
            template = DocxTemplate(root / "template.docx")
            template.render({"name": "ClosyTalk"})
            template.save(root / "document.docx")
            self.assertEqual(Document(root / "document.docx").paragraphs[0].text, "Customer ClosyTalk")

            pd.DataFrame({"Amount": [12, 30]}).to_excel(root / "sheet.xlsx", index=False)
            sheet = load_workbook(root / "sheet.xlsx").active
            self.assertEqual([sheet["A2"].value, sheet["A3"].value], [12, 30])

            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(3), Inches(1))
            shape.text = "Native diagram"
            deck.save(root / "deck.pptx")
            self.assertEqual(Presentation(root / "deck.pptx").slides[0].shapes[0].text, "Native diagram")

            canvas = Canvas(str(root / "document.pdf"))
            canvas.drawString(72, 720, "Offline document")
            canvas.save()
            self.assertIn("Offline document", run("pdftotext", str(root / "document.pdf"), "-"))
            run("pdftoppm", "-singlefile", "-scale-to", "400", "-png", str(root / "document.pdf"), str(root / "page"))
            with Image.open(root / "page.png") as image:
                self.assertEqual(max(image.size), 400)
                self.assertLess(min(image.convert("L").getextrema()), 200)

            plt.plot([1, 2, 3], [2, 4, 3])
            plt.savefig(root / "chart.png")
            plt.close()
            with Image.open(root / "chart.png") as image:
                self.assertGreater(image.width, 100)
                self.assertLess(min(image.convert("L").getextrema()), 100)

    def test_fonts_and_licenses(self):
        for family in ("Liberation Sans", "Carlito", "DejaVu Sans", "Noto Sans"):
            self.assertIn(family, run("fc-match", "-f", "%{family}", family))
        self.assertGreater(Path("/usr/local/share/licenses/node/LICENSE").stat().st_size, 1000)
        self.assertGreater(Path("/usr/local/share/licenses/agent-browser/LICENSE").stat().st_size, 1000)
        self.assertGreater(Path("/usr/local/share/licenses/ffmpeg/COPYING.LGPLv2.1").stat().st_size, 1000)
        self.assertGreater(Path("/usr/local/share/licenses/libvpx/LICENSE").stat().st_size, 1000)
        self.assertTrue(Path("/usr/share/doc/fonts-crosextra-carlito/copyright").is_file())
        chromium_licenses = list(Path("/opt/chromium").rglob("LICENSE.headless_shell"))
        self.assertEqual(len(chromium_licenses), 1)
        self.assertGreater(chromium_licenses[0].stat().st_size, 1000)
        self.assertEqual(list(Path("/opt/chromium").rglob("*widevine*")), [], "Do not ship the proprietary Widevine CDM")

    def test_bundled_skills_offline(self):
        root = Path(os.environ["AGENT_BROWSER_SKILLS_DIR"])
        self.assertEqual({p.name for p in root.iterdir()}, {"core", "dogfood", "derive-client"})
        for name in ("core", "dogfood", "derive-client"):
            self.assertIn("name:", (root / name / "SKILL.md").read_text())
        self.assertIn("core", run("agent-browser", "skills"))

    def test_browser_form_snapshot_and_screenshot(self):
        from PIL import Image
        html = b'''<!doctype html><html><title>Offline sandbox test</title>
        <body style="background:#e9f2ff;color:#102840;padding:32px;font:24px sans-serif">
        <h1>Offline sandbox test</h1><label>Name <input id="name"></label>
        <button id="save" onclick="document.querySelector('#result').textContent='Saved '+document.querySelector('#name').value">Save</button>
        <p id="result">Waiting</p></body></html>'''

        class Handler(http.server.BaseHTTPRequestHandler):
            def do_GET(self):
                self.send_response(200)
                self.send_header("Content-Type", "text/html")
                self.end_headers()
                self.wfile.write(html)

            def log_message(self, *_):
                pass

        server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()

        def browser(*args):
            response = json.loads(run("agent-browser", "--session", "offline-test", "--json", *args))
            self.assertTrue(response["success"], response)
            return response["data"]

        try:
            browser("open", f"http://127.0.0.1:{server.server_port}")
            snapshot = browser("snapshot", "-i")
            self.assertIn("Name", json.dumps(snapshot))
            self.assertIn("Save", json.dumps(snapshot))
            refs = snapshot["refs"]
            textbox = next(ref for ref, item in refs.items() if item["role"] == "textbox" and item["name"] == "Name")
            button = next(ref for ref, item in refs.items() if item["role"] == "button" and item["name"] == "Save")
            browser("fill", f"@{textbox}", "ClosyTalk")
            browser("click", f"@{button}")
            self.assertIn("Saved ClosyTalk", json.dumps(browser("get", "text", "#result")))
            with tempfile.TemporaryDirectory() as directory:
                shot = Path(directory) / "screenshot.png"
                browser("screenshot", str(shot))
                with Image.open(shot) as image:
                    self.assertGreaterEqual(image.width, 800)
                    self.assertLess(min(image.convert("L").getextrema()), 100)
                    self.assertGreater(max(image.convert("L").getextrema()), 200)
        finally:
            try:
                browser("close")
            finally:
                server.shutdown()
                server.server_close()
                thread.join(timeout=5)


if __name__ == "__main__":
    unittest.main(verbosity=2)
